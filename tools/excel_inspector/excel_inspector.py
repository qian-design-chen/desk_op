"""Core logic for batch inspecting Excel files."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}
SUPPORTED_FORMATS = "、".join(sorted(EXCEL_EXTENSIONS))


@dataclass
class MatchResult:
    file: str
    sheet: str
    cell: str
    value: str


def _norm_text(value: object) -> str:
    return re.sub(r"\s+", "", str(value).lower())


def _to_number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = (
        str(value)
        .strip()
        .replace(",", "")
        .replace("￥", "")
        .replace("¥", "")
        .replace("$", "")
        .replace("%", "")
        .replace(" ", "")
    )
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return float(text)
    except ValueError:
        return None


def scan_excel_files(root: Path, recursive: bool = True) -> list[Path]:
    if not root.is_dir():
        return []
    if recursive:
        candidates = (p for p in root.rglob("*") if p.is_file())
    else:
        candidates = (p for p in root.iterdir() if p.is_file())
    return sorted(
        p for p in candidates if p.suffix.lower() in EXCEL_EXTENSIONS
    )


def iter_cells(path: Path) -> Iterator[tuple[str, str, object]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            yield sheet.title, cell.coordinate, cell.value
        finally:
            workbook.close()
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            for row in range(sheet.nrows):
                for col in range(sheet.ncols):
                    value = sheet.cell_value(row, col)
                    if value != "":
                        yield sheet.name, xlrd.formula.cellname(row, col), value
    else:
        return


def search_in_file(
    path: Path,
    mode: str,
    keyword: str = "",
    min_value: float | None = None,
    max_value: float | None = None,
) -> list[MatchResult]:
    matches: list[MatchResult] = []
    file_name = str(path)
    needle = _norm_text(keyword) if keyword else ""
    for sheet, cell, value in iter_cells(path):
        if mode == "text":
            if needle and needle in _norm_text(value):
                matches.append(MatchResult(file_name, sheet, cell, str(value)))
            continue
        number = _to_number(value)
        if number is None:
            continue
        if mode == "gt" and min_value is not None and number > min_value:
            matches.append(MatchResult(file_name, sheet, cell, str(value)))
        elif mode == "lt" and max_value is not None and number < max_value:
            matches.append(MatchResult(file_name, sheet, cell, str(value)))
        elif (
            mode == "range"
            and min_value is not None
            and max_value is not None
            and min_value <= number <= max_value
        ):
            matches.append(MatchResult(file_name, sheet, cell, str(value)))
    return matches


def inspect_files(
    files: list[Path],
    mode: str,
    keyword: str = "",
    min_value: float | None = None,
    max_value: float | None = None,
) -> tuple[list[MatchResult], list[tuple[str, str]], dict[str, int]]:
    matches: list[MatchResult] = []
    errors: list[tuple[str, str]] = []
    per_file: dict[str, int] = {}
    for path in files:
        try:
            file_matches = search_in_file(path, mode, keyword, min_value, max_value)
        except Exception as exc:
            errors.append((str(path), str(exc)))
            continue
        matches.extend(file_matches)
        per_file[str(path)] = len(file_matches)
    return matches, errors, per_file


def count_locations(matches: list[MatchResult]) -> dict[str, int]:
    result: dict[str, int] = {}
    for match in matches:
        key = f"{match.file}｜{match.sheet}｜{match.cell}"
        result[key] = result.get(key, 0) + 1
    return result
