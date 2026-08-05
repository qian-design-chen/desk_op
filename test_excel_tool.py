"""Smoke tests for the embedded Excel batch inspection tool."""

import shutil
import tempfile
from pathlib import Path

import tkinter as tk

from excel_inspector_tool import open_excel_inspector_gui
from tools.excel_inspector.excel_inspector import (
    inspect_files,
    scan_excel_files,
    search_in_file,
)


def make_sample_excel(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"] = "苹果"
    ws["B1"] = "apple 123"
    ws["A2"] = 5
    ws["B2"] = 15
    ws["C2"] = 25
    ws["A3"] = "香蕉"
    ws["B3"] = "Apple"
    wb.save(path)


def main() -> None:
    root = tk.Tk()
    root.withdraw()
    win = open_excel_inspector_gui(root)
    root.update()
    assert win.winfo_exists()

    temp_dir = Path(tempfile.gettempdir()) / "excel_inspector_test"
    if temp_dir.exists():
        shutil.rmtree(temp_dir)
    temp_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = temp_dir / "sample.xlsx"
    make_sample_excel(xlsx_path)

    files = scan_excel_files(temp_dir, recursive=True)
    assert files == [xlsx_path], files

    text_matches = search_in_file(xlsx_path, "text", keyword="苹果")
    assert len(text_matches) == 1, text_matches
    fuzzy_matches = search_in_file(xlsx_path, "text", keyword="app")
    assert len(fuzzy_matches) == 2, fuzzy_matches

    gt_matches = search_in_file(xlsx_path, "gt", min_value=10)
    assert len(gt_matches) == 2, gt_matches
    lt_matches = search_in_file(xlsx_path, "lt", max_value=10)
    assert len(lt_matches) == 1, lt_matches
    range_matches = search_in_file(xlsx_path, "range", min_value=10, max_value=20)
    assert len(range_matches) == 1, range_matches

    matches, errors, per_file = inspect_files(
        [xlsx_path],
        "text",
        keyword="apple",
    )
    assert len(matches) == 2
    assert not errors
    assert per_file[str(xlsx_path)] == 2

    win.destroy()
    root.destroy()
    shutil.rmtree(temp_dir, ignore_errors=True)
    print("excel tool smoke test OK")


if __name__ == "__main__":
    main()
